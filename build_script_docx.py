"""Build DOCX version of the (shortened) speaker script."""

import docx
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x49, 0x7D)
RED = RGBColor(0xC0, 0x50, 0x4D)
GREEN = RGBColor(0x3B, 0x7A, 0x3B)
GRAY = RGBColor(0x59, 0x59, 0x59)
ORANGE = RGBColor(0xE8, 0x6A, 0x17)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
LIGHT = 'F2F4F8'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def p(text='', size=11, bold=False, italic=False, color=BLACK, align=None,
      space_after=4, space_before=0):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    if align:
        par.alignment = align
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    return par


def heading(text):
    par = p(text, size=13.5, bold=True, color=NAVY, space_before=12, space_after=4)
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), '1F497D')
    pbdr.append(bottom)
    pPr.append(pbdr)
    return par


def shade(par, color=LIGHT):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    pPr.append(shd)


def speech(lines):
    """lines: list of (text, dict) run tuples per paragraph."""
    for para in lines:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.left_indent = Cm(0.35)
        shade(par)
        for text, fmt in para:
            r = par.add_run(text)
            r.font.size = Pt(fmt.get('size', 10.5))
            r.font.bold = fmt.get('bold', False)
            r.font.italic = fmt.get('italic', False)
            r.font.color.rgb = fmt.get('color', BLACK)
            r.font.name = 'Calibri'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def qa(q_text, a_text):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(5)
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run('Q: ' + q_text)
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RED
    par2 = doc.add_paragraph()
    par2.paragraph_format.space_after = Pt(5)
    par2.paragraph_format.left_indent = Cm(0.35)
    r2 = par2.add_run('A: ' + a_text)
    r2.font.size = Pt(10.5)


# ================= TITLE =================
p('DIAMOND 演讲逐页讲稿（精简版）', size=18, bold=True, color=NAVY,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
p('Speaker Script (short) — World Models and Video Consistency', size=10.5, color=GRAY,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
p('约 1,025 词 ≈ 6.8–7.9 分钟（130–150 wpm），在 10 分钟时限内留出缓冲 · 配套 12 页 PPT',
  size=9.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

# ================= SLIDE 1 =================
heading('Slide 1 — Title（0:00–0:20）')
speech([[('Good morning. Today: world models, and the one property that decides whether they '
          'work — ', {}),
         ('video consistency', {'bold': True}),
         ('. The paper is ', {}),
         ('Diffusion for World Modeling: Visual Details Matter in Atari', {'italic': True}),
         (' — DIAMOND — from NeurIPS 2024. It was the first work to make diffusion usable as a '
          'world model for reinforcement learning, and it did so by fixing a consistency problem. '
          'I\'ll show you the problem, the fix, and the numbers.', {})]])

# ================= SLIDE 2 =================
heading('Slide 2 — Motivation（0:20–1:10）')
speech([
    [('Two problems. ', {}),
     ('One:', {'bold': True}),
     (' RL is sample-hungry — a classic Atari agent needs ~50 million steps. World models fix '
      'this: learn the dynamics from data, then train the agent ', {}),
     ('in imagination', {'italic': True}),
     (' with only ~100k real steps, about two hours of play.', {})],
    [('Two:', {'bold': True}),
     (' current world models — IRIS, DreamerV3, STORM — encode frames as ', {}),
     ('discrete tokens', {'italic': True}),
     ('. Quantization keeps rollouts stable, but throws away few-pixel details. Enemies '
      'flicker, rewards vanish, scores drift — and these errors compound. The agent trains on '
      'corrupted video.', {})],
    [('The bridge:', {'bold': True}),
     (' diffusion operates in a continuous space, no quantization bottleneck. The paper\'s '
      'question: can a diffusion model ', {}),
     ('be', {'italic': True}),
     (' the world model — and does keeping the visual details actually make the agent better? '
      'Both answers: yes.', {})],
])

# ================= SLIDE 3 =================
heading('Slide 3 — DIAMOND: A Diffusion Model as the Environment（1:10–1:55）')
speech([
    [('This is the full system — the paper\'s Figure 1: at each step, a diffusion model '
      'denoises the next frame, conditioned on past observations and actions.', {})],
    [('Three components. ', {}),
     ('The world model:', {'bold': True}),
     (' a 2D U-Net built on EDM. Past frames are concatenated channel-wise to the noisy next '
      'frame; actions enter through adaptive group-normalization. EDM\'s noise schedule is '
      'what makes sampling stable in just one to three steps.', {})],
    [('A reward/termination model:', {'bold': True}),
     (' a small CNN-LSTM predicting scalar reward and episode end — it handles partial '
      'observability.', {})],
    [('The agent:', {'bold': True}),
     (' an actor-critic, shared CNN-LSTM, trained with REINFORCE plus a value baseline, '
      'entirely inside imagination.', {})],
    [('The loop is standard: collect real data, retrain the world model, train the agent '
      'inside it, repeat.', {})],
])

# ================= SLIDE 4 =================
heading('Slide 4 — Design Choice 1: EDM Beats DDPM（1:55–2:35）')
speech([
    [('First design choice: which diffusion framework? DDPM looks natural — but on the left, '
      'its trajectories corrupt: colors shift, error compounds.', {})],
    [('Why: DDPM predicts the ', {}),
     ('noise', {'italic': True}),
     ('. When noise dominates, the optimal prediction collapses to the identity map — the '
      'model has a terrible score estimate at the start of sampling, and the error '
      'accumulates over rollout.', {})],
    [('EDM adaptively mixes signal and noise: when noise dominates, the network predicts the ', {}),
     ('clean image', {'italic': True}),
     (' — a good score estimate even with almost no signal. On the right: stable, even with '
      'one denoising step.', {})],
    [('First consistency lesson: ', {}),
     ('the diffusion formulation itself determines how fast imagined video drifts.', {'bold': True})],
])

# ================= SLIDE 5 =================
heading('Slide 5 — Design Choice 2: 3 Steps, Not 1（2:35–3:15）')
speech([
    [('But one step isn\'t always enough. Boxing: the opponent\'s moves are unpredictable — '
      'partial observability. With one step, the model predicts the ', {}),
     ('expectation', {'italic': True}),
     (' over possible next frames; when that distribution is multimodal, the expectation is '
      'an out-of-distribution blur — the top row.', {})],
    [('Three iterative steps drive sampling into one specific mode — crisp and consistent, '
      'bottom row. The ablation confirms it: on the ten best games, mean human-normalized '
      'score drops from 3.05 with three steps to 1.96 with one.', {})],
    [('Note the tension: each step is a full network pass, and a world model renders '
      'thousands of frames per episode. DIAMOND\'s achievement is being stable enough that '
      'three steps suffice.', {})],
])

# ================= SLIDE 6 =================
heading('Slide 6 — The Core Result: Consistent Videos（3:15–4:00）')
speech([
    [('Now the paper\'s central comparison: the same imagined trajectories from IRIS and '
      'DIAMOND.', {})],
    [('Watch the white boxes — IRIS. Asterix: an enemy becomes a reward and back. Breakout: '
      'bricks and score drift. Road Runner: rewards flicker in and out.', {})],
    [('None of this happens with DIAMOND. And it\'s ', {}),
     ('not', {'italic': True}),
     (' bought with compute: same 64×64 resolution, three network evaluations per frame '
      'versus sixteen for IRIS, fewer parameters.', {})],
    [('Why it matters: every imagined frame is a training sample. An agent that should '
      'chase rewards and avoid enemies is, in IRIS, literally trained on frames where the '
      'two are swapped. Few-pixel inconsistency is ', {}),
     ('label noise for the policy', {'bold': True}),
     ('.', {})],
])

# ================= SLIDE 7 =================
heading('Slide 7 — Atari 100k Results（4:00–4:45）')
speech([
    [('The benchmark: 26 games, 100k environment steps each — two hours of play — five '
      'seeds.', {})],
    [('DIAMOND reaches a ', {}),
     ('mean human-normalized score of 1.46', {'bold': True}),
     (': superhuman on average, a new best among agents trained entirely inside a world '
      'model, and superhuman on 11 of 26 games.', {})],
    [('The wins concentrate where tiny details decide: Asterix 3,698 versus STORM\'s 1,028; '
      'Breakout 132 versus IRIS\'s 84; Road Runner 20,673 — best on the board.', {})],
    [('The fair comparison matters: the agent is essentially the same as prior work\'s, so '
      'the gap traces to the world model — to the consistency of the video it imagines.', {})],
])

# ================= SLIDE 8 =================
heading('Slide 8 — Stability, Measured（4:45–5:10）')
speech([
    [('Consistency, quantified: one thousand imagined steps of Breakout from the same start '
      'frame, same expert actions, averaged over four hundred reference trajectories.', {})],
    [('DDPM drifts more and more — even ten steps stay worse than one-step EDM. And EDM\'s '
      'residual drift is ', {}),
     ('trajectory difference', {'italic': True}),
     (' — an equally valid but different play — not the color corruption of DDPM.', {})],
    [('The point: consistency is not aesthetics. Inconsistency is noise on the agent\'s '
      'training signal, and it compounds with rollout length. DIAMOND attacks it at the '
      'source.', {})],
])

# ================= SLIDE 9 =================
heading('Slide 9 — Scaling Up: CS:GO / Driving（5:10–5:45）')
speech([
    [('Does the recipe survive reality? They trained the same diffusion world model — no RL, '
      'no engine — on static CS:GO footage and on motorway driving.', {})],
    [('People could ', {}),
     ('play inside it live', {'italic': True}),
     (', keyboard and mouse, stable over minutes of interaction.', {})],
    [('Numbers: FVD ', {}),
     ('34.8', {'bold': True}),
     (' on CS:GO versus 85.7 for the best IRIS and 509 for DreamerV3. Driving: 80.3 versus '
      '277 and 734.', {})],
    [('The same recipe that kept Asterix consistent keeps a first-person 3D world '
      'consistent.', {})],
])

# ================= SLIDE 10 =================
heading('Slide 10 — Limitations（5:45–6:30）')
speech([
    [('Being honest about the limits: 64×64 resolution; sampling at about 10 FPS; the ', {}),
     ('agent', {'italic': True}),
     (' is sample-efficient but the ', {}),
     ('world model', {'italic': True}),
     (' is not — about one GPU-year for the full benchmark; one world model per game.', {})],
    [('Most importantly: visual consistency is ', {}),
     ('not', {'italic': True}),
     (' physical consistency. Rollouts look right but are not physically grounded, and '
      'long-horizon drift remains — an active frontier, and the direction of my own ongoing '
      'project.', {})],
])

# ================= SLIDE 11 =================
heading('Slide 11 — Takeaways（6:30–7:10）')
speech([
    [('Five takeaways. One: ', {}),
     ('consistency is the bottleneck', {'bold': True}),
     (' — discrete tokens lose the few-pixel details the reward signal depends on. Two: ', {}),
     ('diffusion fixes it', {'bold': True}),
     (' — continuous, no quantization bottleneck, stable rollouts. Three: ', {}),
     ('consistency becomes performance', {'bold': True}),
     (' — 1.46, superhuman on 11 of 26. Four: ', {}),
     ('it scales', {'bold': True}),
     (' — CS:GO FVD 34.8 with interactive human play. Five: ', {}),
     ('the frontier is open', {'bold': True}),
     (' — resolution, speed, sample efficiency, physical grounding.', {})],
])

# ================= SLIDE 12 =================
heading('Slide 12 — Thank You（7:10–7:40）')
speech([[('Thank you. Code, agents, and playable world models are released at the project '
          'page. Happy to take questions.', {})]])

# ================= BACKUP =================
heading('备用问答 Backup Q&A')
qa('Why not just a bigger discrete VAE, like IRIS with 64 tokens?',
   'More tokens help a bit (IRIS K=64 in the 3D experiments), but it doesn\'t close the gap — '
   'DIAMOND still achieves far lower FVD and LPIPS at comparable parameters. The issue is the '
   'discrete bottleneck itself, not just its size.')
qa('Isn\'t 3 NFE per frame very expensive compared to a transformer step?',
   'Yes, and the paper is explicit that the world model — not the agent — is the '
   'compute-hungry part (~10 FPS on Atari; 1.03 GPU-years total). Subsequent latent diffusion '
   'world models address exactly this.')
qa('Does this work with stochastic environments?',
   'Yes — diffusion is a natural fit: the denoising process samples from the predictive '
   'distribution. Boxing is precisely a partially observable, stochastic setting where '
   'multi-step sampling resolves multimodality.')
qa('Why Atari — isn\'t it solved?',
   'Atari 100k is the standard sample-efficiency benchmark — 2 hours of gameplay instead of '
   '50M steps. It\'s the right testbed for world-model agents, and the visual-detail failure '
   'mode is cleanest to diagnose there. CS:GO/driving show the recipe scaling to real 3D '
   'footage.')

doc.save('diamond-speaker-script.docx')
print('saved diamond-speaker-script.docx')
